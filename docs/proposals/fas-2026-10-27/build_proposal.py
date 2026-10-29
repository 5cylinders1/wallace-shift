"""Build the editable FAS proposal from its reviewed Markdown source."""
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[3]
SOURCE = Path(__file__).with_name("proposal.md")
DEST = ROOT / "output/docx/wallace-shift-fas-proposal-revised-2026-10-27.docx"
LOGO = ROOT / "output/assets/wallace-shift-logo-black.png"


def normalize(s):
    return s.replace("–", "-").replace("—", "-").replace("\u2011", "-")


def inline(p, text, bold=False, size=None):
    for token in re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\([^)]+\))', normalize(text)):
        if not token:
            continue
        match = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)
        if match:
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), p.part.relate_to(match[2], RT.HYPERLINK, is_external=True))
            r = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            c = OxmlElement('w:color'); c.set(qn('w:val'), '000000'); props.append(c)
            u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); props.append(u)
            r.append(props)
            t = OxmlElement('w:t'); t.text = match[1]; r.append(t)
            link.append(r); p._p.append(link)
        else:
            is_bold = token.startswith('**') and token.endswith('**')
            run = p.add_run(token[2:-2] if is_bold else token)
            run.bold = True if bold or is_bold else None
            if size: run.font.size = Pt(size)


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(.75)
sec.top_margin = Inches(.7)
sec.bottom_margin = Inches(.65)
sec.header_distance = Inches(.3)
sec.footer_distance = Inches(.3)

for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Header', 'Footer'):
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Arial')
    style._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Arial')
    for color in style._element.xpath('./w:rPr/w:color'):
        for attr in ('themeColor', 'themeTint', 'themeShade'):
            color.attrib.pop(qn('w:' + attr), None)
    for border in style._element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)

normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.1
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.widow_control = True

for name, size, before, after in [('Title', 25, 0, 12), ('Heading 1', 21, 0, 13), ('Heading 2', 12.5, 11, 6)]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.font.bold = name != 'Title'
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

hp = sec.header.paragraphs[0]
hp.paragraph_format.tab_stops.clear_all()
hp.paragraph_format.tab_stops.add_tab_stop(Inches(7), WD_TAB_ALIGNMENT.RIGHT)
hp.paragraph_format.space_after = Pt(0)
rr = hp.add_run('WALLACE SHIFT LLC\tFAS WEBSITE PROPOSAL')
rr.font.size = Pt(8)

fp = sec.footer.paragraphs[0]
doc.styles['Footer'].font.size = Pt(8)
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.paragraph_format.space_after = Pt(0)
run = fp.add_run('Page '); run.font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); fp._p.append(field)


def table(rows):
    n = len(rows[0])
    widths = {3: [1.45, 1.55, 4.0], 4: [3.05, .95, 1.5, 1.5], 6: [1.35, 1.08, 1.08, 1.13, 1.06, 1.3], 7: [1.0, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0], 2:[4.4,2.6]}[n]
    if rows[0][0] == 'FAS approval milestone': widths = [4.4, 1.0, 1.6]
    t = doc.add_table(rows=0, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col,w in zip(t.columns,widths): col.width = Inches(w)
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
    t._tbl.tblPr.append(borders)
    for i, row in enumerate(rows):
        cells=t.add_row().cells
        trpr=t.rows[-1]._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
        if i == 0:
            rep=OxmlElement('w:tblHeader'); trpr.append(rep)
        for j,(cell,value,w) in enumerate(zip(cells,row,widths)):
            cell.width=Inches(w)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'333333' if i==0 else ('F3F4F5' if i%2==0 else 'FFFFFF'));tcpr.append(fill)
            margins=OxmlElement('w:tcMar')
            for side, val in [('top',90),('bottom',90),('left',85),('right',85)]:
                m=OxmlElement('w:'+side);m.set(qn('w:w'),str(val));m.set(qn('w:type'),'dxa');margins.append(m)
            tcpr.append(margins)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.06
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or (n==3 and rows[0][0]=='Phase') else WD_ALIGN_PARAGRAPH.CENTER
            inline(p,value,bold=(i==0 or row[0]=='Total'),size=9.3 if n==6 else 10)
            for run in p.runs:
                run.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(1)
    p.paragraph_format.space_before=Pt(0)
    p.paragraph_format.line_spacing=Pt(3)
    p.add_run().font.size=Pt(3)


lines=SOURCE.read_text().splitlines()
i=0
titles=0
next_page=False
while i<len(lines):
    line=lines[i].strip()
    if not line:
        i+=1; continue
    if line=='<!-- pagebreak -->':
        next_page=True;i+=1;continue
    signature = re.fullmatch(r'<!-- signature: (.+) -->', line)
    if signature:
        # Embed the reviewed signature cutout with its alpha channel intact.
        if doc.paragraphs:
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.keep_with_next = True
        picture = p.add_run().add_picture(str(ROOT / signature[1]), width=Inches(2.5))
        picture._inline.docPr.set('descr', 'Matthew Wallace signature with photographed paper background removed')
        i += 1
        continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',x) for x in cells): rows.append(cells)
            i+=1
        table(rows);continue
    if line.startswith('# '):
        if titles == 0:
            brand = doc.add_paragraph()
            brand.paragraph_format.keep_with_next = True
            brand.paragraph_format.space_after = Pt(4)
            picture = brand.add_run().add_picture(str(LOGO), width=Inches(.8))
            picture._inline.docPr.set('descr', 'Wallace Shift logo')
        p=doc.add_paragraph(style='Title' if titles==0 else 'Heading 1')
        if next_page:
            p.paragraph_format.page_break_before=True
            next_page=False
        inline(p,line[2:]);titles+=1
    elif line.startswith('## '):
        p=doc.add_paragraph(style='Heading 2');inline(p,line[3:])
    else:
        p=doc.add_paragraph();inline(p,line)
    i+=1

doc.core_properties.title='FAS website development proposal'
doc.core_properties.subject='Revised October 27 2026 following FAS scope update'
doc.core_properties.author='Matthew Wallace | Wallace Shift LLC'
doc.core_properties.keywords='FAS, WordPress, website development, proposal'
doc.core_properties.comments=''
DEST.parent.mkdir(parents=True,exist_ok=True)
doc.save(DEST)
print(DEST)
